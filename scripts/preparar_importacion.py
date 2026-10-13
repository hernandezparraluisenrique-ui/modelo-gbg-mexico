
import csv
from pathlib import Path
from datetime import date
from openpyxl import load_workbook

# Rutas de los archivos
carpeta_proyecto = Path(__file__).resolve().parent.parent
archivo_excel = carpeta_proyecto / "Libro1.xlsx"
archivo_csv = carpeta_proyecto / "datos_importacion.csv"

if not archivo_excel.exists():
    raise FileNotFoundError(
        f"No se encontró el Excel: {archivo_excel}"
    )

libro = load_workbook(archivo_excel, data_only=True)
hoja = libro["Hoja1"]

# Los meses están en las columnas C a Y.
# Cada mes se convierte en el primer día de ese mes.
meses_espanol = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4,
    "may": 5, "jun": 6, "jul": 7, "ago": 8,
    "sep": 9, "oct": 10, "nov": 11, "dic": 12
}

# Filas donde comienzan y terminan los estados.
bloques = [
    ("Bovino", 4, 35),
    ("Porcino", 41, 72)
]

encabezados = [
    "nombre_estado",
    "hospedador",
    "periodo",
    "valor_original",
    "valor_numerico",
    "estado_registro"
]

registros = []
estados_encontrados = set()

for hospedador, fila_inicio, fila_fin in bloques:
    for fila in range(fila_inicio, fila_fin + 1):
        nombre_estado = hoja.cell(fila, 2).value

        if not nombre_estado:
            continue

        nombre_estado = str(nombre_estado).strip()
        estados_encontrados.add(nombre_estado)

        # Columnas C a Y: 23 meses
        for columna in range(3, 26):
            encabezado_mes = hoja.cell(3, columna).value

            # El segundo bloque usa los mismos meses,
            # pero también verificamos su encabezado.
            if hospedador == "Porcino":
                encabezado_mes = hoja.cell(40, columna).value

            texto_mes = str(encabezado_mes).strip().lower()
            abreviatura, anio_corto = texto_mes.split("-")
            mes = meses_espanol[abreviatura]
            anio = 2000 + int(anio_corto)

            periodo = date(anio, mes, 1).isoformat()
            valor = hoja.cell(fila, columna).value

            if valor is None:
                # Celda vacía: no significa cero.
                valor_original = ""
                valor_numerico = ""
                estado_registro = "vacio"

            elif isinstance(valor, str) and valor.strip() == "-":
                valor_original = "-"
                valor_numerico = ""
                estado_registro = "guion"

            elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
                if valor < 0 or not float(valor).is_integer():
                    raise ValueError(
                        f"Valor numérico inesperado en "
                        f"{nombre_estado}, {hospedador}, "
                        f"{periodo}: {valor!r}"
                    )

                valor_numerico = int(valor)
                valor_original = str(valor_numerico)
                estado_registro = "numerico"

            else:
                raise ValueError(
                    f"Valor inesperado en "
                    f"{nombre_estado}, {hospedador}, "
                    f"{periodo}: {valor!r}"
                )

            registros.append([
                nombre_estado,
                hospedador,
                periodo,
                valor_original,
                valor_numerico,
                estado_registro
            ])

# Crear el CSV sin modificar el Excel original.
with open(archivo_csv, "w", newline="", encoding="utf-8-sig") as archivo:
    escritor = csv.writer(archivo)
    escritor.writerow(encabezados)
    escritor.writerows(registros)

# Resumen de control
print("Preparación terminada.")
print(f"Archivo creado: {archivo_csv}")
print(f"Estados distintos: {len(estados_encontrados)}")
print(f"Registros preparados: {len(registros)}")

for tipo in ("numerico", "guion", "vacio"):
    cantidad = sum(1 for r in registros if r[5] == tipo)
    print(f"Registros tipo '{tipo}': {cantidad}")

print("\nPrimeros 5 registros:")
for registro in registros[:5]:
    print(registro)