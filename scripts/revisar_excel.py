
from openpyxl import load_workbook
from pathlib import Path

archivo = Path(__file__).resolve().parent.parent / "Libro1.xlsx"

if not archivo.exists():
    raise FileNotFoundError(f"No se encontró el archivo: {archivo}")

libro = load_workbook(archivo, data_only=True)

for hoja in libro.worksheets:
    print(f"\n{'=' * 50}")
    print(f"HOJA: {hoja.title}")
    print(f"Filas: {hoja.max_row} | Columnas: {hoja.max_column}")
    print("=" * 50)

    for numero_fila in range(1, hoja.max_row + 1):
        valor_a = hoja.cell(numero_fila, 1).value
        valor_b = hoja.cell(numero_fila, 2).value

        if valor_a is not None or valor_b is not None:
            print(
                f"Fila {numero_fila}: "
                f"A={valor_a!r} | B={valor_b!r}"
            )

            if str(valor_b).strip().lower() == "estado":
                meses = [
                    hoja.cell(numero_fila, columna).value
                    for columna in range(3, hoja.max_column + 1)
                ]
                print("MESES:", meses)